from typing import Any, Dict, List, Optional, Tuple, Union
import io
import frappe
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


@frappe.whitelist()
def download_voucher_excel(doctype: str = "Petty Cash Entry", docname: str = None, voucher_name: str = None):
    """
    Unified entry point for downloading voucher Excel workbooks directly from desk action buttons.
    Supports both docname & voucher_name params.
    """
    target_name = docname or voucher_name
    if not target_name:
        frappe.throw("Voucher name / Docname is required for Excel export.", frappe.ValidationError)

    if doctype == "Petty Cash Entry":
        return export_petty_cash_excel(voucher_name=target_name)
    else:
        return export_petty_cash_excel(voucher_name=target_name)


@frappe.whitelist()
def export_petty_cash_excel(voucher_name: str = None, docname: str = None, doctype: str = None):
    """Generates and downloads a beautifully styled Excel workbook for a Petty Cash Voucher."""
    target_name = voucher_name or docname
    if not target_name:
        frappe.throw("Voucher name is required for Excel export.", frappe.ValidationError)

    if not frappe.has_permission("Petty Cash Entry", "read", target_name):
        frappe.throw("You do not have permission to export this Petty Cash Voucher.", frappe.PermissionError)

    doc = frappe.get_doc("Petty Cash Entry", target_name)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Voucher {doc.name}"
    ws.views.sheetView[0].showGridLines = True

    # Palette
    title_font = Font(name="Calibri", size=15, bold=True, color="1E3A8A")
    sec_font = Font(name="Calibri", size=11, bold=True, color="1E293B")
    th_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    th_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    meta_label_font = Font(name="Calibri", size=10, bold=True, color="475569")
    meta_val_font = Font(name="Calibri", size=10, color="0F172A")
    total_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Masked Account Format for Security
    raw_acc = str(getattr(doc, "custodian_bank_account", "") or "")
    if len(raw_acc) >= 4 and not raw_acc.startswith("••••"):
        masked_acc = f"•••• •••• •••• {raw_acc[-4:]}"
    else:
        masked_acc = raw_acc or "N/A"

    # 1. Title Banner
    ws.merge_cells("A1:L1")
    ws["A1"] = f"PETTY CASH VOUCHER STATEMENT — {doc.name}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    # 2. Metadata Section
    meta = [
        ("Voucher ID:", doc.name, "Posting Date:", str(doc.posting_date or "")),
        ("Claim Title:", getattr(doc, "claim_title", None) or "N/A", "Workflow Status:", doc.status or "Draft"),
        ("Company:", doc.company or "N/A", "Payment Batch ID:", getattr(doc, "batch_id", None) or "N/A"),
        ("JV Number (Accounts):", getattr(doc, "jv_number", None) or "Pending Accounts Posting", "Total Amount (₹):", f"₹ {float(doc.total_amount or 0):,.2f}"),
        ("Beneficiary Employee:", getattr(doc, "custodian", None) or "N/A", "Beneficiary Name:", getattr(doc, "beneficiary_name", None) or "N/A"),
        ("Bank Name:", getattr(doc, "bank_name", None) or "N/A", "Bank A/C No:", masked_acc),
        ("Bank IFSC Code:", getattr(doc, "custodian_ifsc_code", None) or "N/A", "Generated On:", str(frappe.utils.today()))
    ]

    r = 3
    for m in meta:
        ws.cell(row=r, column=1, value=m[0]).font = meta_label_font
        ws.cell(row=r, column=2, value=m[1]).font = meta_val_font
        ws.cell(row=r, column=4, value=m[2]).font = meta_label_font
        ws.cell(row=r, column=5, value=m[3]).font = meta_val_font
        r += 1

    r += 1
    # 3. Line Items Table (Updated Columns Matching New Standard)
    ws.cell(row=r, column=1, value="EXPENSE LINE ITEMS").font = sec_font
    r += 1

    headers = [
        "S.NO",
        "DATE",
        "BRANCH",
        "EXPENSES OF HEAD",
        "PARTICULARS (Merchant / Details)",
        "GST?",
        "MERCHANT GSTIN",
        "BILL # / TXN ID",
        "AMOUNT (₹)",
        "STAFF NAME",
        "RECEIPT",
        "REMARK",
        "DECISION"
    ]
    
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=c_idx, value=h)
        cell.font = th_font
        cell.fill = th_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[r].height = 24

    r += 1
    start_line_r = r
    lines = doc.expense_lines or []
    category_summary = {}
    for idx, line in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=str(getattr(line, "expense_date", None) or getattr(doc, "posting_date", "") or "")).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=3, value=getattr(line, "branch", "") or "Head Office").alignment = Alignment(horizontal="center")
        
        is_before_accounts = doc.status in ("Draft", "Pending Admin L1", "Pending Admin L2", "Rejected")
        raw_head = getattr(line, "expense_head", None)
        if raw_head == "OTHER / VENDOR LEDGER" and getattr(line, "custom_ledger_name", None):
            exp_head = getattr(line, "custom_ledger_name", "OTHER / VENDOR LEDGER")
        elif raw_head:
            exp_head = raw_head
        elif is_before_accounts:
            exp_head = "Pending Accounts Classification"
        else:
            exp_head = getattr(line, "expense_category", "") or "OFFICE EXPENSES"
        ws.cell(row=r, column=4, value=exp_head)
        
        ws.cell(row=r, column=5, value=getattr(line, "merchant_name", "") or "")
        
        # GST?
        is_gst = int(getattr(line, "is_gst", 0) or 0)
        ws.cell(row=r, column=6, value="Yes" if is_gst else "No").alignment = Alignment(horizontal="center")
        
        # Merchant GSTIN
        ws.cell(row=r, column=7, value=getattr(line, "merchant_gstin", "") or "N/A").alignment = Alignment(horizontal="center")
        
        # Bill # / Txn ID
        ws.cell(row=r, column=8, value=getattr(line, "bill_number", "") or getattr(line, "transaction_id", "") or "N/A").alignment = Alignment(horizontal="center")
        
        # Amount (₹) - Column 9 (I)
        line_amt = float(getattr(line, "amount", 0.0) or 0.0)
        amt_cell = ws.cell(row=r, column=9, value=line_amt)
        amt_cell.number_format = '₹ #,##0.00'
        amt_cell.alignment = Alignment(horizontal="right")
        
        # Accumulate Category for Pivot Summary
        category_summary[exp_head] = category_summary.get(exp_head, 0.0) + line_amt
        
        # Staff Name
        staff_disp = getattr(line, "staff_name", "") or getattr(line, "employee", "") or ""
        ws.cell(row=r, column=10, value=staff_disp).alignment = Alignment(horizontal="center")
        
        # Receipt
        receipt_url = getattr(line, "receipt_attachment", None) or ""
        receipt_cell = ws.cell(row=r, column=11, value="View Receipt" if receipt_url else "No Receipt")
        if receipt_url:
            receipt_cell.hyperlink = receipt_url
            receipt_cell.font = Font(color="2563EB", underline="single")
        receipt_cell.alignment = Alignment(horizontal="center")

        # Remarks / Bill type
        ws.cell(row=r, column=12, value=getattr(line, "remarks", "") or getattr(line, "description", "") or "BILL")

        # Decision
        if getattr(line, "is_disputed", 0):
            d_val = f"⚠️ Disputed: {getattr(line, 'dispute_reason', '') or 'Disputed'}"
            d_cell = ws.cell(row=r, column=13, value=d_val)
            d_cell.font = Font(name="Calibri", size=10, bold=True, color="B91C1C")
        else:
            d_cell = ws.cell(row=r, column=13, value="✅ OK / Approved")
            d_cell.font = Font(name="Calibri", size=10, color="15803D")
        d_cell.alignment = Alignment(horizontal="center")

        for c in range(1, 14):
            ws.cell(row=r, column=c).border = thin_border
        r += 1

    if lines:
        # Total Summary Row
        ws.cell(row=r, column=1, value="TOTAL").font = total_font
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=1).fill = total_fill
        for c in range(2, 9):
            ws.cell(row=r, column=c).fill = total_fill
        
        # Formula for Amount (Column I)
        tot_cell = ws.cell(row=r, column=9, value=f"=SUM(I{start_line_r}:I{r-1})")
        tot_cell.font = total_font
        tot_cell.fill = total_fill
        tot_cell.number_format = '₹ #,##0.00'
        tot_cell.alignment = Alignment(horizontal="right")
        
        for c in range(10, 14):
            ws.cell(row=r, column=c).fill = total_fill
        for c in range(1, 14):
            ws.cell(row=r, column=c).border = thin_border
        r += 2

        # --------------------------------------------------------------------------------------
        # ACCOUNTS SUMMARY / PIVOT BLOCK (Exact Match with Real-World Excel Sheet)
        # --------------------------------------------------------------------------------------
        pivot_title_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        pivot_head_fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
        
        ws.cell(row=r, column=2, value="ACCOUNTS LEDGER SUMMARY (EXPENSES OF HEAD)").font = sec_font
        r += 1
        
        # Pivot Headers: Row Labels | Sum of AMOUNT | JV Number | Date
        jv_val = getattr(doc, "jv_number", None) or "JV Pending"
        post_date_str = str(doc.posting_date or frappe.utils.today())
        
        ws.cell(row=r, column=2, value="Row Labels").font = Font(name="Calibri", size=10, bold=True, color="0F172A")
        ws.cell(row=r, column=2).fill = pivot_head_fill
        ws.cell(row=r, column=2).border = thin_border
        
        ws.cell(row=r, column=3, value="Sum of AMOUNT").font = Font(name="Calibri", size=10, bold=True, color="0F172A")
        ws.cell(row=r, column=3).fill = pivot_head_fill
        ws.cell(row=r, column=3).alignment = Alignment(horizontal="right")
        ws.cell(row=r, column=3).border = thin_border
        
        ws.cell(row=r, column=4, value=jv_val).font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")
        ws.cell(row=r, column=4).fill = pivot_head_fill
        ws.cell(row=r, column=4).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=4).border = thin_border
        
        ws.cell(row=r, column=5, value=post_date_str).font = Font(name="Calibri", size=10, bold=True, color="0F172A")
        ws.cell(row=r, column=5).fill = pivot_head_fill
        ws.cell(row=r, column=5).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=5).border = thin_border
        r += 1
        
        pivot_start_r = r
        for cat_label, cat_amt in sorted(category_summary.items()):
            ws.cell(row=r, column=2, value=cat_label).border = thin_border
            p_cell = ws.cell(row=r, column=3, value=cat_amt)
            p_cell.number_format = '₹ #,##0.00'
            p_cell.alignment = Alignment(horizontal="right")
            p_cell.border = thin_border
            ws.cell(row=r, column=4, value="").border = thin_border
            ws.cell(row=r, column=5, value="").border = thin_border
            r += 1
            
        # Grand Total of Pivot
        ws.cell(row=r, column=2, value="Grand Total").font = Font(name="Calibri", size=10, bold=True, color="000000")
        ws.cell(row=r, column=2).fill = total_fill
        ws.cell(row=r, column=2).border = thin_border
        
        p_tot_cell = ws.cell(row=r, column=3, value=f"=SUM(C{pivot_start_r}:C{r-1})")
        p_tot_cell.font = Font(name="Calibri", size=10, bold=True, color="000000")
        p_tot_cell.fill = total_fill
        p_tot_cell.number_format = '₹ #,##0.00'
        p_tot_cell.alignment = Alignment(horizontal="right")
        p_tot_cell.border = thin_border
        
        ws.cell(row=r, column=4, value="").fill = total_fill
        ws.cell(row=r, column=4).border = thin_border
        ws.cell(row=r, column=5, value="").fill = total_fill
        ws.cell(row=r, column=5).border = thin_border
        r += 2

    # 4. Approval Audit Trail
    if getattr(doc, "approval_trail", None) and len(doc.approval_trail) > 0:
        ws.cell(row=r, column=1, value="APPROVAL AUDIT TRAIL").font = sec_font
        r += 1
        trail_headers = ["Level", "Action", "Action Taken By / Approver", "Timestamp", "Remarks / Notes"]
        for c_idx, h in enumerate(trail_headers, 1):
            cell = ws.cell(row=r, column=c_idx, value=h)
            cell.font = th_font
            cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        r += 1
        for tr in doc.approval_trail:
            lvl_str = f"Level {tr.level_number}" if getattr(tr, "level_number", None) else (getattr(tr, "level_name", None) or "")
            user_str = getattr(tr, "action_taken_by", None) or getattr(tr, "designated_approver", None) or ""
            action_str = getattr(tr, "action", None) or ""
            time_str = str(getattr(tr, "action_timestamp", None) or "")
            rem_str = getattr(tr, "remarks", None) or ""

            ws.cell(row=r, column=1, value=lvl_str).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=2, value=action_str).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=3, value=user_str)
            ws.cell(row=r, column=4, value=time_str).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=5, value=rem_str)
            for c in range(1, 6):
                ws.cell(row=r, column=c).border = thin_border
            r += 1

    # Adjust widths
    for col in ws.columns:
        vals = [len(str(cell.value or '')) for cell in col if cell.row > 1]
        max_len = max(vals) if vals else 10
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 38), 12)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    frappe.response['filename'] = f"Petty_Cash_{doc.name}.xlsx"
    frappe.response['filecontent'] = out.getvalue()
    frappe.response['type'] = 'binary'

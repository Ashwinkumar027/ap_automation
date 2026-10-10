# -*- coding: utf-8 -*-
# Copyright (c) 2026, Quantique and contributors
# For license information, please see license.txt

import frappe
from frappe.utils import getdate, flt, cstr, formatdate, nowdate
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import io
import zipfile
import json
import re
from collections import OrderedDict
from typing import Any, Dict, List, Optional, Tuple, Union

def generate_reimbursement_excel_bytes(voucher_name: str):
    """
    Generates Excel bytes and standard filename for a single reimbursement claim voucher.
    """
    if not voucher_name:
        frappe.throw("Voucher name is required.")

    doc = frappe.get_doc("Employee Reimbursement Claim", voucher_name)

    # Fetch employee bank details if not present on doc
    emp_id = doc.employee
    emp = frappe.db.get_value("Employee", emp_id, ["employee_name", "bank_name", "bank_ac_no", "ifsc_code", "department", "company"], as_dict=True) or {}
    
    emp_name = doc.get("employee_name") or emp.get("employee_name") or emp_id
    id_number = emp_id or ""
    bank_name = doc.get("bank_name") or emp.get("bank_name") or ""
    ac_number = doc.get("bank_account_number") or emp.get("bank_ac_no") or ""
    ifsc_code = doc.get("bank_ifsc_code") or emp.get("ifsc_code") or ""
    name_as_per_bank = emp_name
    company_name = doc.get("company") or emp.get("company") or "Quanticus Software Solutions Private Limited"

    category = doc.get("claim_category") or "General Expense"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reimbursement Voucher"
    ws.views.sheetView[0].showGridLines = True

    # Font styles
    font_company_title = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    font_doc_title = Font(name="Calibri", size=11, bold=True, color="475569")
    font_header_bold = Font(name="Calibri", size=10, bold=True, color="000000")
    font_regular = Font(name="Calibri", size=10, bold=False, color="000000")
    font_notes_bold = Font(name="Calibri", size=9, bold=True, color="000000")
    font_notes = Font(name="Calibri", size=9, bold=False, color="334155")
    
    # Fills
    fill_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_section = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    
    # Alignments
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)
    
    # Borders
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(bottom=Side(border_style="medium", color="475569"), left=thin_border_side, right=thin_border_side, top=thin_border_side)

    max_col = 6
    if category == "Dinner Allowance":
        max_col = 7

    # 1. Company Name & Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    title_cell = ws.cell(row=1, column=1, value=company_name)
    title_cell.font = font_company_title
    title_cell.alignment = align_center

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    sub_title = ws.cell(row=2, column=1, value=f"REIMBURSEMENT VOUCHER - {category.upper()} (Voucher #{voucher_name})")
    sub_title.font = font_doc_title
    sub_title.alignment = align_center

    curr_row = 4

    # 2. Employee & Bank Meta Grid
    meta_left = [
        ("ID Number", id_number),
        ("Name", emp_name),
        ("Department", emp.get("department") or doc.get("department") or "General"),
        ("Claim Category", category),
        ("Claim Date", formatdate(doc.posting_date, "dd-mm-yyyy") if doc.posting_date else "")
    ]

    meta_right = [
        ("Bank Name", bank_name),
        ("A/C Number", str(ac_number)),
        ("IFSC Code", ifsc_code),
        ("Name as per Bank", name_as_per_bank),
        ("Status", doc.docstatus == 1 and "Approved / Audited" or doc.status or "Submitted")
    ]

    for i in range(max(len(meta_left), len(meta_right))):
        ws.row_dimensions[curr_row].height = 20
        if i < len(meta_left):
            k, v = meta_left[i]
            c1 = ws.cell(row=curr_row, column=1, value=k)
            c1.font = font_header_bold
            c1.fill = fill_header
            c1.border = thin_border
            
            c2 = ws.cell(row=curr_row, column=2, value=v)
            c2.font = font_regular
            c2.border = thin_border
            if max_col >= 6:
                ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=3)
                ws.cell(row=curr_row, column=3).border = thin_border

        if i < len(meta_right):
            k, v = meta_right[i]
            col_k = 4 if max_col >= 6 else 3
            col_v = 5 if max_col >= 6 else 4
            
            c3 = ws.cell(row=curr_row, column=col_k, value=k)
            c3.font = font_header_bold
            c3.fill = fill_header
            c3.border = thin_border
            
            c4 = ws.cell(row=curr_row, column=col_v, value=v)
            c4.font = font_regular
            c4.border = thin_border
            if max_col >= 6 and col_v < max_col:
                ws.merge_cells(start_row=curr_row, start_column=col_v, end_row=curr_row, end_column=max_col)
                for cx in range(col_v + 1, max_col + 1):
                    ws.cell(row=curr_row, column=cx).border = thin_border

        curr_row += 1

    curr_row += 1

    # 3. Itemized Tables based on Category
    if category == "Client Visit Travel":
        # Table Headers: S.No | Date | Mode of Travel | From -> To | Distance (KM) | Amount (INR)
        headers = ["S.No", "Visit Date", "Mode of Travel", "Route (From -> To)", "Distance (KM)", "Total Leg Amount (INR)"]
        ws.row_dimensions[curr_row].height = 24
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=curr_row, column=col_idx, value=h)
            cell.font = font_header_bold
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border

        curr_row += 1
        start_data_row = curr_row

        legs = doc.get("client_visit_legs") or []
        for idx, leg in enumerate(legs, 1):
            ws.row_dimensions[curr_row].height = 20
            ws.cell(row=curr_row, column=1, value=idx).alignment = align_center
            ws.cell(row=curr_row, column=2, value=formatdate(leg.visit_date, "dd-mm-yyyy") if leg.visit_date else "").alignment = align_center
            mode_val = getattr(leg, 'mode_of_travel', None) or getattr(leg, 'travel_mode', None) or 'Travel'
            ws.cell(row=curr_row, column=3, value=mode_val).alignment = align_center
            ws.cell(row=curr_row, column=4, value=f"{leg.from_location or ''} -> {leg.to_location or ''}").alignment = align_left
            ws.cell(row=curr_row, column=5, value=flt(leg.distance_km or 0.0)).alignment = align_right
            amt_cell = ws.cell(row=curr_row, column=6, value=flt(leg.leg_amount or 0.0) + flt(leg.toll_parking_amount or 0.0))
            amt_cell.alignment = align_right
            amt_cell.number_format = '₹ #,##0.00'

            for c in range(1, 7):
                ws.cell(row=curr_row, column=c).border = thin_border
                ws.cell(row=curr_row, column=c).font = font_regular
            curr_row += 1

        end_data_row = curr_row - 1
        if not legs:
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=6)
            ws.cell(row=curr_row, column=1, value="No travel legs recorded").alignment = align_center
            for c in range(1, 7):
                ws.cell(row=curr_row, column=c).border = thin_border
            curr_row += 1
            end_data_row = curr_row - 1

        # Total Row
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=5)
        tot_label = ws.cell(row=curr_row, column=1, value="GRAND TOTAL")
        tot_label.font = font_header_bold
        tot_label.alignment = align_center
        tot_amt = ws.cell(row=curr_row, column=6, value=f"=SUM(F{start_data_row}:F{end_data_row})" if legs else flt(doc.total_claim_amount or 0.0))
        tot_amt.font = font_header_bold
        tot_amt.alignment = align_right
        tot_amt.number_format = '₹ #,##0.00'

        for c in range(1, 7):
            ws.cell(row=curr_row, column=c).border = thick_bottom
            ws.cell(row=curr_row, column=c).fill = fill_section

    elif category == "Dinner Allowance":
        headers = ["S.No", "Expense Date", "Swipe Out Time", "Work Summary", "Eligible Amount (INR)", "Claimed (INR)", "Remarks"]
        ws.row_dimensions[curr_row].height = 24
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=curr_row, column=col_idx, value=h)
            cell.font = font_header_bold
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border

        curr_row += 1
        start_data_row = curr_row
        lines = doc.get("expense_lines") or []
        for idx, line in enumerate(lines, 1):
            ws.row_dimensions[curr_row].height = 20
            ws.cell(row=curr_row, column=1, value=idx).alignment = align_center
            ws.cell(row=curr_row, column=2, value=formatdate(line.expense_date, "dd-mm-yyyy") if line.expense_date else "").alignment = align_center
            ws.cell(row=curr_row, column=3, value=str(line.out_time or "")).alignment = align_center
            ws.cell(row=curr_row, column=4, value=line.description or "Overtime Dinner").alignment = align_left
            
            c5 = ws.cell(row=curr_row, column=5, value=flt(line.amount or 0.0))
            c5.alignment = align_right
            c5.number_format = '₹ #,##0.00'

            c6 = ws.cell(row=curr_row, column=6, value=flt(line.amount or 0.0))
            c6.alignment = align_right
            c6.number_format = '₹ #,##0.00'

            ws.cell(row=curr_row, column=7, value=line.payment_mode_used or "Auto-calculated").alignment = align_left

            for c in range(1, 8):
                ws.cell(row=curr_row, column=c).border = thin_border
                ws.cell(row=curr_row, column=c).font = font_regular
            curr_row += 1

        end_data_row = curr_row - 1
        if not lines:
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=7)
            ws.cell(row=curr_row, column=1, value="No dinner allowance lines").alignment = align_center
            for c in range(1, 8):
                ws.cell(row=curr_row, column=c).border = thin_border
            curr_row += 1
            end_data_row = curr_row - 1

        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=5)
        tot_label = ws.cell(row=curr_row, column=1, value="GRAND TOTAL")
        tot_label.font = font_header_bold
        tot_label.alignment = align_center
        tot_amt = ws.cell(row=curr_row, column=6, value=f"=SUM(F{start_data_row}:F{end_data_row})" if lines else flt(doc.total_claim_amount or 0.0))
        tot_amt.font = font_header_bold
        tot_amt.alignment = align_right
        tot_amt.number_format = '₹ #,##0.00'

        for c in range(1, 8):
            ws.cell(row=curr_row, column=c).border = thick_bottom
            ws.cell(row=curr_row, column=c).fill = fill_section

    else:
        # General Expense, Team Lunch, Training, etc.
        headers = ["S.No", "Expense Date", "Expense Type / Purpose", "Description", "Amount (INR)", "Payment Mode / Remarks"]
        ws.row_dimensions[curr_row].height = 24
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=curr_row, column=col_idx, value=h)
            cell.font = font_header_bold
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border

        curr_row += 1
        start_data_row = curr_row
        lines = doc.get("expense_lines") or []
        for idx, line in enumerate(lines, 1):
            ws.row_dimensions[curr_row].height = 20
            ws.cell(row=curr_row, column=1, value=idx).alignment = align_center
            ws.cell(row=curr_row, column=2, value=formatdate(line.expense_date, "dd-mm-yyyy") if line.expense_date else "").alignment = align_center
            ws.cell(row=curr_row, column=3, value=line.expense_type or category).alignment = align_left
            ws.cell(row=curr_row, column=4, value=line.description or "").alignment = align_left
            
            c5 = ws.cell(row=curr_row, column=5, value=flt(line.amount or 0.0))
            c5.alignment = align_right
            c5.number_format = '₹ #,##0.00'

            ws.cell(row=curr_row, column=6, value=line.payment_mode_used or "").alignment = align_left

            for c in range(1, 7):
                ws.cell(row=curr_row, column=c).border = thin_border
                ws.cell(row=curr_row, column=c).font = font_regular
            curr_row += 1

        end_data_row = curr_row - 1
        if not lines:
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=6)
            ws.cell(row=curr_row, column=1, value="No individual line items recorded").alignment = align_center
            for c in range(1, 7):
                ws.cell(row=curr_row, column=c).border = thin_border
            curr_row += 1
            end_data_row = curr_row - 1

        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=4)
        tot_label = ws.cell(row=curr_row, column=1, value="GRAND TOTAL")
        tot_label.font = font_header_bold
        tot_label.alignment = align_center
        tot_amt = ws.cell(row=curr_row, column=5, value=f"=SUM(E{start_data_row}:E{end_data_row})" if lines else flt(doc.total_claim_amount or 0.0))
        tot_amt.font = font_header_bold
        tot_amt.alignment = align_right
        tot_amt.number_format = '₹ #,##0.00'

        for c in range(1, 7):
            ws.cell(row=curr_row, column=c).border = thick_bottom
            ws.cell(row=curr_row, column=c).fill = fill_section

    curr_row += 2

    # 4. Important Notes
    ws.cell(row=curr_row, column=1, value="Important Notes:").font = font_notes_bold
    curr_row += 1

    notes = [
        "• This voucher is electronically generated from Quanticus AP Automation Accounts Audit Workstation.",
        "• Disbursal will be credited directly to the employee registered bank account above.",
        "• All original invoices and attached proofs have been audited and archived digitally.",
        "• Strict verification of IFSC code and Bank Account Number has been performed."
    ]

    for note in notes:
        ws.cell(row=curr_row, column=1, value=note).font = font_notes
        curr_row += 1

    # Set Column Widths
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 24
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 22
    if category == "Dinner Allowance":
        ws.column_dimensions['G'].width = 24

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Reimbursement_Voucher_{voucher_name}_{category.replace(' ', '_')}.xlsx"
    return filename, output.getvalue()


def _parse_claim_names_list(claim_names: Any) -> List[str]:
    """
    Robust claim names parser that cleanly handles:
    - Lists of strings or integers
    - JSON-encoded arrays
    - Comma-separated strings
    - Nested JSON lists
    """
    if not claim_names:
        return []
    
    if isinstance(claim_names, (list, tuple, set)):
        result = []
        for item in claim_names:
            if isinstance(item, str) and (item.startswith("[") or item.startswith("{")):
                try:
                    parsed = json.loads(item)
                    if isinstance(parsed, list):
                        result.extend([str(x).strip() for x in parsed if x])
                        continue
                except Exception:
                    pass
            if item:
                result.append(str(item).strip())
        return [r for r in result if r]

    if isinstance(claim_names, str):
        claim_names = claim_names.strip()
        if claim_names.startswith("[") and claim_names.endswith("]"):
            try:
                parsed = json.loads(claim_names)
                if isinstance(parsed, list):
                    return [str(x).strip() for x in parsed if x]
            except Exception:
                pass
        
        # Regex split on comma, quotes, brackets
        parts = re.split(r'[,\s"\[\]]+', claim_names)
        return [p.strip() for p in parts if p.strip()]

    return [str(claim_names).strip()]


def generate_consolidated_employee_excel_bytes(claim_names: Any, employee_name: Optional[str] = None, is_batch: bool = False, batch_name: Optional[str] = None) -> Tuple[str, bytes]:
    """
    Generates a consolidated Excel workbook containing all itemized lines across all selected claims.
    
    Modes:
    1. Full Batch Export (is_batch=True or batch_name provided or employee_name starts with 'Full_Batch_' or multiple employees):
       - Sheet 1 ('Itemized Batch Expenses'): Full 13-column itemized breakdown identifying EVERY employee and their bank details on each row.
       - Sheet 2 ('Bank Payout Summary'): Dedicated bank upload/disbursal summary grouped per employee with total payout amounts.
    2. Single Employee Consolidated Export:
       - Single consolidated worksheet with employee header and all vouchers/expense lines.
    """
    raw_claims = _parse_claim_names_list(claim_names)
    if not raw_claims:
        frappe.throw("No claims found to generate consolidated Excel.")

    # Load all docs
    docs = []
    for c_name in raw_claims:
        if not c_name:
            continue
        try:
            d = frappe.get_doc("Employee Reimbursement Claim", c_name)
            docs.append(d)
        except Exception as e:
            frappe.log_error(f"Error loading claim {c_name} for consolidated Excel: {e}")

    if not docs:
        frappe.throw("None of the specified claims could be loaded.")

    first_doc = docs[0]
    
    # Check if this is a multi-employee batch
    distinct_emps = list(OrderedDict.fromkeys([d.employee for d in docs if d.employee]))
    is_multi_employee_batch = is_batch or bool(batch_name) or (employee_name and employee_name.startswith("Full_Batch_")) or (len(distinct_emps) > 1)

    # Global company legal entity
    company_name = first_doc.get("company")
    if not company_name and first_doc.employee:
        company_name = frappe.db.get_value("Employee", first_doc.employee, "company")
    company_name = company_name or "Quanticus Software Solutions Private Limited"

    wb = openpyxl.Workbook()
    
    # Styles
    font_company_title = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    font_doc_title = Font(name="Calibri", size=11, bold=True, color="475569")
    font_header_bold = Font(name="Calibri", size=10, bold=True, color="000000")
    font_regular = Font(name="Calibri", size=10, bold=False, color="000000")
    font_notes_bold = Font(name="Calibri", size=9, bold=True, color="000000")
    font_notes = Font(name="Calibri", size=9, bold=False, color="334155")
    
    fill_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_section = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_accent = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
    fill_highlight = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(bottom=Side(border_style="medium", color="475569"), left=thin_border_side, right=thin_border_side, top=thin_border_side)

    # Pre-fetch all employee info for fast mapping
    emp_map = {}
    for emp_id in distinct_emps:
        info = frappe.db.get_value("Employee", emp_id, ["employee_name", "bank_name", "bank_ac_no", "ifsc_code", "department", "company"], as_dict=True) or {}
        emp_map[emp_id] = info

    # Flatten all line items across all vouchers
    all_lines = []
    global_sno = 1

    for doc in docs:
        v_no = doc.name
        cat = doc.get("claim_category") or "General Expense"
        doc_emp_id = doc.employee
        emp_info = emp_map.get(doc_emp_id, {})
        
        doc_emp_name = doc.get("employee_name") or emp_info.get("employee_name") or doc_emp_id
        doc_bank_name = doc.get("bank_name") or emp_info.get("bank_name") or ""
        doc_ac_no = str(doc.get("bank_account_number") or emp_info.get("bank_ac_no") or "").strip()
        doc_ifsc = str(doc.get("bank_ifsc_code") or emp_info.get("ifsc_code") or "").strip()
        doc_company = doc.get("company") or emp_info.get("company") or company_name

        if cat == "Client Visit Travel":
            legs = doc.get("client_visit_legs") or []
            if legs:
                for leg in legs:
                    dt_str = formatdate(leg.visit_date, "dd-mm-yyyy") if leg.visit_date else ""
                    mode = getattr(leg, 'mode_of_travel', None) or getattr(leg, 'travel_mode', None) or 'Travel'
                    from_loc = getattr(leg, 'from_location', '') or ''
                    to_loc = getattr(leg, 'to_location', '') or ''
                    dist = getattr(leg, 'distance_km', 0) or 0
                    desc = f"{mode}: {from_loc} -> {to_loc} ({dist} km)"
                    amt = flt(leg.leg_amount or 0.0) + flt(leg.toll_parking_amount or 0.0)
                    rem = f"Mileage + Tolls | {leg.client_name or ''}"
                    all_lines.append({
                        "sno": global_sno,
                        "voucher": v_no,
                        "emp_id": doc_emp_id,
                        "emp_name": doc_emp_name,
                        "bank_name": doc_bank_name,
                        "bank_ac_no": doc_ac_no,
                        "ifsc_code": doc_ifsc,
                        "category": cat,
                        "date": dt_str,
                        "desc": desc,
                        "amount": amt,
                        "remarks": rem,
                        "company": doc_company
                    })
                    global_sno += 1
            else:
                all_lines.append({
                    "sno": global_sno,
                    "voucher": v_no,
                    "emp_id": doc_emp_id,
                    "emp_name": doc_emp_name,
                    "bank_name": doc_bank_name,
                    "bank_ac_no": doc_ac_no,
                    "ifsc_code": doc_ifsc,
                    "category": cat,
                    "date": formatdate(doc.posting_date, "dd-mm-yyyy") if doc.posting_date else "",
                    "desc": f"Client Travel: {doc.client_name or ''}",
                    "amount": flt(doc.total_claim_amount or 0.0),
                    "remarks": doc.remarks or "",
                    "company": doc_company
                })
                global_sno += 1
        else:
            lines = doc.get("expense_lines") or []
            if lines:
                for line in lines:
                    dt_str = formatdate(line.expense_date, "dd-mm-yyyy") if line.expense_date else ""
                    desc = line.description or line.expense_type or "Expense"
                    amt = flt(line.amount or 0.0)
                    rem = line.payment_mode_used or ""
                    if line.out_time:
                        rem += f" (Out: {line.out_time})"
                    all_lines.append({
                        "sno": global_sno,
                        "voucher": v_no,
                        "emp_id": doc_emp_id,
                        "emp_name": doc_emp_name,
                        "bank_name": doc_bank_name,
                        "bank_ac_no": doc_ac_no,
                        "ifsc_code": doc_ifsc,
                        "category": cat,
                        "date": dt_str,
                        "desc": desc,
                        "amount": amt,
                        "remarks": rem,
                        "company": doc_company
                    })
                    global_sno += 1
            else:
                all_lines.append({
                    "sno": global_sno,
                    "voucher": v_no,
                    "emp_id": doc_emp_id,
                    "emp_name": doc_emp_name,
                    "bank_name": doc_bank_name,
                    "bank_ac_no": doc_ac_no,
                    "ifsc_code": doc_ifsc,
                    "category": cat,
                    "date": formatdate(doc.posting_date, "dd-mm-yyyy") if doc.posting_date else "",
                    "desc": cat,
                    "amount": flt(doc.total_claim_amount or 0.0),
                    "remarks": doc.remarks or "",
                    "company": doc_company
                })
                global_sno += 1

    total_batch_sum = sum(l["amount"] for l in all_lines)

    # ---------------------------------------------------------
    # CASE 1: MULTI-EMPLOYEE BATCH CONSOLIDATED EXCEL
    # ---------------------------------------------------------
    if is_multi_employee_batch:
        batch_label = batch_name or (employee_name.replace("Full_Batch_", "") if employee_name and employee_name.startswith("Full_Batch_") else "Weekly Batch")
        
        # TAB 1: Itemized Batch Expenses
        ws1 = wb.active
        ws1.title = "Itemized Batch Expenses"
        ws1.views.sheetView[0].showGridLines = True
        max_col1 = 13

        # Title
        ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col1)
        t1 = ws1.cell(row=1, column=1, value=company_name)
        t1.font = font_company_title
        t1.alignment = align_center

        ws1.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col1)
        sub1 = ws1.cell(row=2, column=1, value=f"WEEKLY ACCOUNTS AUDIT - FULL BATCH ITEMIZE VOUCHER ({batch_label} - {len(docs)} Claims across {len(distinct_emps)} Employees)")
        sub1.font = font_doc_title
        sub1.alignment = align_center

        curr_row = 4
        
        # Batch Summary Metadata Box
        meta_left = [
            ("Batch Reference", batch_label),
            ("Total Claims Count", f"{len(docs)} Vouchers"),
            ("Total Employees Count", f"{len(distinct_emps)} Employees"),
            ("Export Date", formatdate(nowdate(), "dd-mm-yyyy"))
        ]
        meta_right = [
            ("Total Batch Amount", f"₹ {total_batch_sum:,.2f}"),
            ("Batch Audit Status", "Verified / Audited for Payment"),
            ("Disbursal Mode", "Direct Bank Transfer (NEFT / RTGS / IMPS)"),
            ("Company Legal Entity", company_name)
        ]

        for i in range(max(len(meta_left), len(meta_right))):
            ws1.row_dimensions[curr_row].height = 20
            if i < len(meta_left):
                k, v = meta_left[i]
                c1 = ws1.cell(row=curr_row, column=1, value=k)
                c1.font = font_header_bold
                c1.fill = fill_header
                c1.border = thin_border

                ws1.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=5)
                c2 = ws1.cell(row=curr_row, column=2, value=v)
                c2.font = font_regular
                c2.border = thin_border
                for cx in range(3, 6):
                    ws1.cell(row=curr_row, column=cx).border = thin_border

            if i < len(meta_right):
                k, v = meta_right[i]
                c3 = ws1.cell(row=curr_row, column=6, value=k)
                c3.font = font_header_bold
                c3.fill = fill_header
                c3.border = thin_border

                ws1.merge_cells(start_row=curr_row, start_column=7, end_row=curr_row, end_column=max_col1)
                c4 = ws1.cell(row=curr_row, column=7, value=v)
                c4.font = font_regular
                c4.border = thin_border
                for cx in range(8, max_col1 + 1):
                    ws1.cell(row=curr_row, column=cx).border = thin_border

            curr_row += 1

        curr_row += 1

        # Table Header (13 Columns)
        headers1 = [
            "S.No", "Voucher #", "Employee ID", "Employee Name", "Bank Name", 
            "Bank A/C Number", "IFSC Code", "Category", "Expense Date", 
            "Description / Purpose", "Amount (INR)", "Payment Mode / Remarks", "Company Legal Entity"
        ]
        ws1.row_dimensions[curr_row].height = 26
        for c_idx, h_text in enumerate(headers1, 1):
            cell = ws1.cell(row=curr_row, column=c_idx, value=h_text)
            cell.font = font_header_bold
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border

        curr_row += 1
        row_data_start = curr_row

        for line in all_lines:
            ws1.row_dimensions[curr_row].height = 20
            ws1.cell(row=curr_row, column=1, value=line["sno"]).alignment = align_center
            
            v_cell = ws1.cell(row=curr_row, column=2, value=line["voucher"])
            v_cell.alignment = align_center
            v_cell.font = Font(name="Calibri", size=10, bold=True, color="2563EB")

            ws1.cell(row=curr_row, column=3, value=line["emp_id"]).alignment = align_center
            ws1.cell(row=curr_row, column=4, value=line["emp_name"]).alignment = align_left
            ws1.cell(row=curr_row, column=5, value=line["bank_name"]).alignment = align_center
            
            ac_c = ws1.cell(row=curr_row, column=6, value=str(line["bank_ac_no"]))
            ac_c.alignment = align_center
            ac_c.number_format = '@'

            ws1.cell(row=curr_row, column=7, value=line["ifsc_code"]).alignment = align_center
            ws1.cell(row=curr_row, column=8, value=line["category"]).alignment = align_center
            ws1.cell(row=curr_row, column=9, value=line["date"]).alignment = align_center
            ws1.cell(row=curr_row, column=10, value=line["desc"]).alignment = align_left

            amt_cell = ws1.cell(row=curr_row, column=11, value=line["amount"])
            amt_cell.alignment = align_right
            amt_cell.number_format = '₹ #,##0.00'

            ws1.cell(row=curr_row, column=12, value=line["remarks"]).alignment = align_left
            ws1.cell(row=curr_row, column=13, value=line["company"]).alignment = align_center

            for c_idx in range(1, max_col1 + 1):
                ws1.cell(row=curr_row, column=c_idx).border = thin_border
                if c_idx != 2:
                    ws1.cell(row=curr_row, column=c_idx).font = font_regular

            curr_row += 1

        row_data_end = curr_row - 1

        # Grand Total Row
        ws1.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=10)
        tot_lbl = ws1.cell(row=curr_row, column=1, value=f"GRAND TOTAL ({len(all_lines)} Rows Across {len(docs)} Vouchers for {len(distinct_emps)} Employees)")
        tot_lbl.font = font_header_bold
        tot_lbl.alignment = align_center

        tot_amt_c = ws1.cell(row=curr_row, column=11, value=f"=SUM(K{row_data_start}:K{row_data_end})")
        tot_amt_c.font = font_header_bold
        tot_amt_c.alignment = align_right
        tot_amt_c.number_format = '₹ #,##0.00'

        for c_idx in range(1, max_col1 + 1):
            ws1.cell(row=curr_row, column=c_idx).border = thick_bottom
            ws1.cell(row=curr_row, column=c_idx).fill = fill_section

        curr_row += 2

        # Notes Footer
        ws1.cell(row=curr_row, column=1, value="Important Accounts & Disbursal Notes:").font = font_notes_bold
        curr_row += 1
        notes = [
            "• Consolidated batch report generated from Quanticus AP Automation Accounts Audit Workstation.",
            "• Payments will strictly be released to each employee's verified registered bank account (Third-party payments prohibited).",
            "• All attached invoices, tax bills, and digital audit proofs are archived and verified in the system.",
            "• Accounts team must cross-check the Account Number and IFSC Code before initiating bank batch release."
        ]
        for n in notes:
            ws1.cell(row=curr_row, column=1, value=n).font = font_notes
            curr_row += 1

        # Column widths for Sheet 1
        ws1.column_dimensions['A'].width = 8   # S.No
        ws1.column_dimensions['B'].width = 20  # Voucher #
        ws1.column_dimensions['C'].width = 14  # Emp ID
        ws1.column_dimensions['D'].width = 24  # Emp Name
        ws1.column_dimensions['E'].width = 18  # Bank Name
        ws1.column_dimensions['F'].width = 20  # A/C No
        ws1.column_dimensions['G'].width = 15  # IFSC
        ws1.column_dimensions['H'].width = 20  # Category
        ws1.column_dimensions['I'].width = 14  # Date
        ws1.column_dimensions['J'].width = 40  # Description
        ws1.column_dimensions['K'].width = 18  # Amount
        ws1.column_dimensions['L'].width = 28  # Remarks
        ws1.column_dimensions['M'].width = 32  # Company

        # ---------------------------------------------------------
        # TAB 2: Bank Payout Summary (Grouped by Employee)
        # ---------------------------------------------------------
        ws2 = wb.create_sheet(title="Bank Payout Summary")
        ws2.views.sheetView[0].showGridLines = True
        max_col2 = 8

        # Title
        ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col2)
        t2 = ws2.cell(row=1, column=1, value=company_name)
        t2.font = font_company_title
        t2.alignment = align_center

        ws2.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col2)
        sub2 = ws2.cell(row=2, column=1, value=f"EMPLOYEE BANK DISBURSAL & PAYOUT SUMMARY - {batch_label}")
        sub2.font = font_doc_title
        sub2.alignment = align_center

        s2_row = 4
        # Metadata
        ws2.merge_cells(start_row=s2_row, start_column=1, end_row=s2_row, end_column=2)
        ws2.cell(row=s2_row, column=1, value="Batch Reference").font = font_header_bold
        ws2.cell(row=s2_row, column=1).fill = fill_header
        ws2.cell(row=s2_row, column=1).border = thin_border
        ws2.cell(row=s2_row, column=2).border = thin_border

        ws2.merge_cells(start_row=s2_row, start_column=3, end_row=s2_row, end_column=4)
        ws2.cell(row=s2_row, column=3, value=batch_label).font = font_regular
        ws2.cell(row=s2_row, column=3).border = thin_border
        ws2.cell(row=s2_row, column=4).border = thin_border

        ws2.merge_cells(start_row=s2_row, start_column=5, end_row=s2_row, end_column=6)
        ws2.cell(row=s2_row, column=5, value="Total Payout Amount").font = font_header_bold
        ws2.cell(row=s2_row, column=5).fill = fill_header
        ws2.cell(row=s2_row, column=5).border = thin_border
        ws2.cell(row=s2_row, column=6).border = thin_border

        ws2.merge_cells(start_row=s2_row, start_column=7, end_row=s2_row, end_column=8)
        c_tot = ws2.cell(row=s2_row, column=7, value=f"₹ {total_batch_sum:,.2f}")
        c_tot.font = font_header_bold
        c_tot.border = thin_border
        ws2.cell(row=s2_row, column=8).border = thin_border

        s2_row += 2

        # Summary Table Headers
        headers2 = [
            "S.No", "Employee ID", "Employee Name / Beneficiary", "Bank Name", 
            "Bank Account Number", "IFSC Code", "Vouchers Count", "Total Payout Amount (INR)"
        ]
        ws2.row_dimensions[s2_row].height = 26
        for c_idx, h_text in enumerate(headers2, 1):
            cell = ws2.cell(row=s2_row, column=c_idx, value=h_text)
            cell.font = font_header_bold
            cell.fill = fill_accent
            cell.alignment = align_center
            cell.border = thin_border

        s2_row += 1
        s2_data_start = s2_row

        # Group lines by Employee
        emp_summary = OrderedDict()
        for l in all_lines:
            e_key = l["emp_id"]
            if e_key not in emp_summary:
                emp_summary[e_key] = {
                    "emp_id": e_key,
                    "emp_name": l["emp_name"],
                    "bank_name": l["bank_name"],
                    "bank_ac_no": l["bank_ac_no"],
                    "ifsc_code": l["ifsc_code"],
                    "vouchers": set(),
                    "total_amount": 0.0
                }
            emp_summary[e_key]["vouchers"].add(l["voucher"])
            emp_summary[e_key]["total_amount"] += flt(l["amount"])

        emp_sno = 1
        for e_id, s_data in emp_summary.items():
            ws2.row_dimensions[s2_row].height = 20
            ws2.cell(row=s2_row, column=1, value=emp_sno).alignment = align_center
            ws2.cell(row=s2_row, column=2, value=s_data["emp_id"]).alignment = align_center
            ws2.cell(row=s2_row, column=3, value=s_data["emp_name"]).alignment = align_left
            ws2.cell(row=s2_row, column=4, value=s_data["bank_name"]).alignment = align_center
            
            ac_cell2 = ws2.cell(row=s2_row, column=5, value=str(s_data["bank_ac_no"]))
            ac_cell2.alignment = align_center
            ac_cell2.number_format = '@'

            ws2.cell(row=s2_row, column=6, value=s_data["ifsc_code"]).alignment = align_center
            ws2.cell(row=s2_row, column=7, value=len(s_data["vouchers"])).alignment = align_center

            payout_cell = ws2.cell(row=s2_row, column=8, value=s_data["total_amount"])
            payout_cell.alignment = align_right
            payout_cell.font = Font(name="Calibri", size=10, bold=True, color="047857")
            payout_cell.number_format = '₹ #,##0.00'

            for c_idx in range(1, max_col2 + 1):
                ws2.cell(row=s2_row, column=c_idx).border = thin_border
                if c_idx != 8:
                    ws2.cell(row=s2_row, column=c_idx).font = font_regular

            s2_row += 1
            emp_sno += 1

        s2_data_end = s2_row - 1

        # Summary Total Row
        ws2.merge_cells(start_row=s2_row, start_column=1, end_row=s2_row, end_column=6)
        tot_s2 = ws2.cell(row=s2_row, column=1, value=f"TOTAL DISBURSAL ({len(emp_summary)} Employees)")
        tot_s2.font = font_header_bold
        tot_s2.alignment = align_center

        v_tot = ws2.cell(row=s2_row, column=7, value=f"=SUM(G{s2_data_start}:G{s2_data_end})")
        v_tot.font = font_header_bold
        v_tot.alignment = align_center

        amt_tot = ws2.cell(row=s2_row, column=8, value=f"=SUM(H{s2_data_start}:H{s2_data_end})")
        amt_tot.font = font_header_bold
        amt_tot.alignment = align_right
        amt_tot.number_format = '₹ #,##0.00'

        for c_idx in range(1, max_col2 + 1):
            ws2.cell(row=s2_row, column=c_idx).border = thick_bottom
            ws2.cell(row=s2_row, column=c_idx).fill = fill_section

        s2_row += 3

        # Disbursal Sign-off Box
        ws2.cell(row=s2_row, column=1, value="Prepared By (Accounts Exe)").font = font_notes_bold
        ws2.cell(row=s2_row, column=4, value="Audited By (Accounts Mgr)").font = font_notes_bold
        ws2.cell(row=s2_row, column=7, value="Approved for Disbursal (Finance Dir)").font = font_notes_bold

        # Column widths for Sheet 2
        ws2.column_dimensions['A'].width = 8   # S.No
        ws2.column_dimensions['B'].width = 16  # Emp ID
        ws2.column_dimensions['C'].width = 28  # Emp Name
        ws2.column_dimensions['D'].width = 20  # Bank Name
        ws2.column_dimensions['E'].width = 24  # Bank A/C
        ws2.column_dimensions['F'].width = 16  # IFSC
        ws2.column_dimensions['G'].width = 16  # Total Vouchers
        ws2.column_dimensions['H'].width = 24  # Total Payout

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"Consolidated_Batch_Reimbursements_{batch_label}.xlsx"
        return filename, output.getvalue()

    # ---------------------------------------------------------
    # CASE 2: SINGLE EMPLOYEE CONSOLIDATED EXCEL
    # ---------------------------------------------------------
    ws = wb.active
    ws.title = "Consolidated Claims"
    ws.views.sheetView[0].showGridLines = True
    max_col = 8

    # Title
    emp_info = emp_map.get(first_doc.employee, {})
    single_emp_name = employee_name or first_doc.get("employee_name") or emp_info.get("employee_name") or first_doc.employee
    single_id = first_doc.employee or ""
    single_bank = first_doc.get("bank_name") or emp_info.get("bank_name") or ""
    single_ac = str(first_doc.get("bank_account_number") or emp_info.get("bank_ac_no") or "")
    single_ifsc = str(first_doc.get("bank_ifsc_code") or emp_info.get("ifsc_code") or "")

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    t = ws.cell(row=1, column=1, value=company_name)
    t.font = font_company_title
    t.alignment = align_center

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    sub = ws.cell(row=2, column=1, value=f"CONSOLIDATED EMPLOYEE REIMBURSEMENT VOUCHER - {single_emp_name.upper()} ({len(docs)} Claims)")
    sub.font = font_doc_title
    sub.alignment = align_center

    curr_row = 4

    meta_left = [
        ("Employee ID", single_id),
        ("Employee Name", single_emp_name),
        ("Department", emp_info.get("department") or first_doc.get("department") or "General"),
        ("Total Claims Count", f"{len(docs)} Vouchers"),
        ("Export Date", formatdate(nowdate(), "dd-mm-yyyy"))
    ]

    meta_right = [
        ("Bank Name", single_bank),
        ("A/C Number", single_ac),
        ("IFSC Code", single_ifsc),
        ("Name as per Bank", single_emp_name),
        ("Batch Status", "Verified / Audited for Payment")
    ]

    for i in range(max(len(meta_left), len(meta_right))):
        ws.row_dimensions[curr_row].height = 20
        if i < len(meta_left):
            k, v = meta_left[i]
            c1 = ws.cell(row=curr_row, column=1, value=k)
            c1.font = font_header_bold
            c1.fill = fill_header
            c1.border = thin_border

            ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=4)
            c2 = ws.cell(row=curr_row, column=2, value=v)
            c2.font = font_regular
            c2.border = thin_border
            ws.cell(row=curr_row, column=3).border = thin_border
            ws.cell(row=curr_row, column=4).border = thin_border

        if i < len(meta_right):
            k, v = meta_right[i]
            c3 = ws.cell(row=curr_row, column=5, value=k)
            c3.font = font_header_bold
            c3.fill = fill_header
            c3.border = thin_border

            ws.merge_cells(start_row=curr_row, start_column=6, end_row=curr_row, end_column=max_col)
            c4 = ws.cell(row=curr_row, column=6, value=v)
            c4.font = font_regular
            c4.border = thin_border
            for col_x in range(7, max_col + 1):
                ws.cell(row=curr_row, column=col_x).border = thin_border

        curr_row += 1

    curr_row += 1

    # Table Header
    headers = ["S.No", "Voucher #", "Category", "Expense Date", "Description / Purpose", "Amount (INR)", "Payment Mode / Remarks", "Company Legal Entity"]
    ws.row_dimensions[curr_row].height = 26
    for col_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(row=curr_row, column=col_idx, value=h_text)
        cell.font = font_header_bold
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    curr_row += 1
    row_data_start = curr_row

    for line in all_lines:
        ws.row_dimensions[curr_row].height = 20
        ws.cell(row=curr_row, column=1, value=line["sno"]).alignment = align_center

        v_cell = ws.cell(row=curr_row, column=2, value=line["voucher"])
        v_cell.alignment = align_center
        v_cell.font = Font(name="Calibri", size=10, bold=True, color="2563EB")

        ws.cell(row=curr_row, column=3, value=line["category"]).alignment = align_center
        ws.cell(row=curr_row, column=4, value=line["date"]).alignment = align_center
        ws.cell(row=curr_row, column=5, value=line["desc"]).alignment = align_left

        amt_cell = ws.cell(row=curr_row, column=6, value=line["amount"])
        amt_cell.alignment = align_right
        amt_cell.number_format = '₹ #,##0.00'

        ws.cell(row=curr_row, column=7, value=line["remarks"]).alignment = align_left
        ws.cell(row=curr_row, column=8, value=line["company"]).alignment = align_center

        for c_idx in range(1, max_col + 1):
            ws.cell(row=curr_row, column=c_idx).border = thin_border
            if c_idx != 2:
                ws.cell(row=curr_row, column=c_idx).font = font_regular

        curr_row += 1

    row_data_end = curr_row - 1

    # Total Row
    ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=5)
    total_label = ws.cell(row=curr_row, column=1, value=f"GRAND TOTAL ({len(all_lines)} Rows Across {len(docs)} Vouchers)")
    total_label.font = font_header_bold
    total_label.alignment = align_center

    total_amt_cell = ws.cell(row=curr_row, column=6, value=f"=SUM(F{row_data_start}:F{row_data_end})")
    total_amt_cell.font = font_header_bold
    total_amt_cell.alignment = align_right
    total_amt_cell.number_format = '₹ #,##0.00'

    for c_idx in range(1, max_col + 1):
        ws.cell(row=curr_row, column=c_idx).border = thick_bottom
        ws.cell(row=curr_row, column=c_idx).fill = fill_section

    curr_row += 2

    # Important Notes
    ws.cell(row=curr_row, column=1, value="Important Notes:").font = font_notes_bold
    curr_row += 1

    notes = [
        "• Consolidated report generated from Quanticus AP Automation Accounts Audit Workstation.",
        "• Payments will only be made to the respective employee's registered bank account; third-party payments are not allowed.",
        "• All attached invoices and digital proofs are audited and archived in the system.",
        "• Crosscheck the account number and IFSC code before disbursal release."
    ]

    for note in notes:
        ws.cell(row=curr_row, column=1, value=note).font = font_notes
        curr_row += 1

    ws.column_dimensions['A'].width = 8   # S.No
    ws.column_dimensions['B'].width = 22  # Voucher #
    ws.column_dimensions['C'].width = 22  # Category
    ws.column_dimensions['D'].width = 14  # Date
    ws.column_dimensions['E'].width = 40  # Description
    ws.column_dimensions['F'].width = 18  # Amount
    ws.column_dimensions['G'].width = 28  # Remarks
    ws.column_dimensions['H'].width = 34  # Company Name

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    clean_emp = "".join(c for c in (single_emp_name or "Employee") if c.isalnum() or c in (" ", "-", "_")).strip()
    filename = f"Consolidated_Reimbursement_Voucher_{clean_emp}.xlsx"
    return filename, output.getvalue()


@frappe.whitelist()
def download_reimbursement_excel(voucher_name: str):
    """
    Dynamically generates the official Employee Expense Reimbursement Voucher
    in Microsoft Excel (.xlsx) format with HRMS legal entity company name.
    """
    filename, filebytes = generate_reimbursement_excel_bytes(voucher_name)
    frappe.response['filename'] = filename
    frappe.response['filecontent'] = filebytes
    frappe.response['type'] = 'binary'


@frappe.whitelist()
def download_consolidated_employee_excel(claim_names: Any = None, employee_name: Optional[str] = None):
    """
    Direct endpoint: Generates ONE single consolidated Excel (.xlsx) containing ALL line items
    across all vouchers for a given employee.
    """
    if isinstance(claim_names, str):
        try:
            claim_names = json.loads(claim_names)
        except Exception:
            claim_names = [claim_names]

    if not claim_names:
        frappe.throw("No voucher names specified for consolidated Excel download.")

    filename, filebytes = generate_consolidated_employee_excel_bytes(claim_names, employee_name=employee_name)
    frappe.response['filename'] = filename
    frappe.response['filecontent'] = filebytes
    frappe.response['type'] = 'binary'


@frappe.whitelist()
def download_batch_consolidated_excel(batch_name: str):
    """
    Direct endpoint: Generates ONE single consolidated Excel (.xlsx) containing ALL claims and line items
    for ALL employees in the entire weekly accounts audit batch, with individual Employee & Bank details
    on every row plus a dedicated Bank Payout Summary tab.
    """
    if not batch_name:
        frappe.throw("Batch name is required.")

    batch = frappe.get_doc("Weekly Accounts Audit Batch", batch_name)
    vouchers = [item.voucher_no for item in (batch.items or []) if item.voucher_no]

    if not vouchers:
        frappe.throw("No vouchers found in this batch to export.")

    filename, filebytes = generate_consolidated_employee_excel_bytes(vouchers, employee_name=f"Full_Batch_{batch_name}", is_batch=True, batch_name=batch_name)
    frappe.response['filename'] = filename
    frappe.response['filecontent'] = filebytes
    frappe.response['type'] = 'binary'


@frappe.whitelist()
def download_multiple_claims_excel_zip(claim_names: Any = None, employee_name: Optional[str] = None):
    """
    Generates Excel vouchers for multiple claims and packages them into a ZIP download.
    If single claim is passed, returns that single .xlsx directly.
    """
    if isinstance(claim_names, str):
        try:
            claim_names = json.loads(claim_names)
        except Exception:
            claim_names = [claim_names]

    if not claim_names:
        frappe.throw("No voucher names specified for Excel download.")

    if len(claim_names) == 1:
        return download_reimbursement_excel(claim_names[0])

    zip_buffer = io.BytesIO()
    file_count = 0
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        clean_claims = _parse_claim_names_list(claim_names)
        for v_name in clean_claims:
            if not v_name:
                continue
            try:
                fn, b_content = generate_reimbursement_excel_bytes(str(v_name))
                zip_file.writestr(fn, b_content)
                file_count += 1
            except Exception as e:
                frappe.log_error(f"Error generating Excel for voucher {v_name}: {e}")

    if file_count == 0:
        frappe.throw("Could not generate Excel vouchers for the selected claims.")

    zip_buffer.seek(0)
    clean_emp = "".join(c for c in (employee_name or "Employee") if c.isalnum() or c in (" ", "-", "_")).strip() or "Employee"
    frappe.response["filename"] = f"{clean_emp}_Individual_Vouchers_Excel.zip"
    frappe.response["filecontent"] = zip_buffer.getvalue()
    frappe.response["type"] = "download"

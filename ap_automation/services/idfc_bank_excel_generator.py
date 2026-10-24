# Copyright (c) 2026, Quanti and contributors
# For license information, please see license.txt

import io
import json
from collections import OrderedDict
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import frappe
from frappe import _
from frappe.utils import flt, formatdate, getdate, nowdate


def get_company_debit_account(company: Optional[str] = None) -> str:
    """
    Retrieves the company IDFC debit account number.
    Checks IDFC Company Credential, Bank Account doctype, with fallback to 10173678610.
    """
    default_acc = "10173678610"
    if not company:
        return default_acc

    try:
        # 1. Check IDFC Company Credential child in AP IDFC Settings
        cred_acc = frappe.db.get_value(
            "IDFC Company Credential",
            {"company": company, "is_active": 1},
            "debit_account_number"
        )
        if cred_acc:
            clean_cred = str(cred_acc).strip()
            # If realistic numeric bank account number
            if clean_cred.isdigit() and len(clean_cred) >= 8:
                return clean_cred

        # 2. Check Company Bank Account
        bank_acc = frappe.db.get_value(
            "Bank Account",
            {"company": company, "is_company_account": 1},
            "bank_account_no"
        )
        if bank_acc:
            clean_bank = str(bank_acc).strip()
            if clean_bank.isdigit() and len(clean_bank) >= 8:
                return clean_bank

    except Exception as e:
        frappe.log_error(f"Error fetching company debit account for {company}: {e}")

    return default_acc


def generate_idfc_bank_excel_bytes(
    payment_rows: List[Dict[str, Any]],
    debit_account_no: str,
    transaction_date_str: Optional[str] = None,
    filename_prefix: str = "IDFC_Bank_Template"
) -> Tuple[str, bytes]:
    """
    Constructs an openpyxl Workbook formatted 100% identically to the official IDFC FIRST Bank
    Bulk Disbursal Upload Template ('IDFC Sheet').
    
    Includes:
    - Row 1: Header (38.25 height, Soft Blue #B4C6E7 fill, 10pt regular Calibri, thin borders)
    - Row 2: Instructions (102.0 height, Pastel Blue #D9E1F2 fill, 10pt italic Calibri, wrapped)
    - Rows 3+: Data rows with exact column widths, '@' text format, and '0.00' amount format.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IDFC Sheet"
    ws.views.sheetView[0].showGridLines = True

    # Column widths matching official template
    col_widths = {
        "A": 22.57,  # Beneficiary Name
        "B": 25.86,  # Beneficiary Account Number
        "C": 21.43,  # IFSC
        "D": 18.14,  # Transaction Type
        "E": 18.57,  # Debit Account Number
        "F": 15.86,  # Transaction Date (DD/MM/YYYY)
        "G": 13.14,  # Amount
        "H": 14.14,  # Currency
        "I": 28.00,  # Beneficiary Email ID
        "J": 23.43,  # Remarks
        "K": 13.00,  # Custom Header – 1
        "L": 13.00,  # Custom Header – 2
        "M": 13.00,  # Custom Header – 3
        "N": 13.00,  # Custom Header – 4
        "O": 13.00,  # Custom Header – 5
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 38.25
    ws.row_dimensions[2].height = 102.0

    headers_r1 = [
        "Beneficiary Name",
        "Beneficiary Account Number",
        "IFSC",
        "Transaction Type",
        "Debit Account Number",
        "Transaction Date",
        "Amount",
        "Currency",
        "Beneficiary Email ID",
        "Remarks",
        "Custom Header – 1",
        "Custom Header – 2",
        "Custom Header – 3",
        "Custom Header – 4",
        "Custom Header – 5",
    ]

    headers_r2 = [
        "Enter beneficiary name.\nMANDATORY",
        "Enter beneficiary account number. \nThis can be IDFC FIRST Bank account or other Bank account.\nMANDATORY",
        "Enter beneficiary bank IFSC code. Required only for Inter bank (NEFT/RTGS) payment.",
        "Enter payment type:\nIFT - Within Bank payment\nNEFT - Inter-Bank(NEFT) payment\nRTGS - Inter-Bank(RTGS) payment\nMANDATORY",
        "Enter debit account number. This should be IDFC FIRST Bank account only. User should have access to do transaction on this account",
        "Enter transaction value date. Should be today's date or future date.\nMANDATORY\nDD/MM/YYYY format",
        "Enter payment amount.\nMANDATORY",
        "Enter transaction currency. Should be INR only.\nMANDATORY",
        "Enter beneficiary email id\nOPTIONAL",
        "Enter remarks\nOPTIONAL",
        "Credit Advice:\nEnter Custom Info -1\nNote: Header label is editable in Row 1\nOPTIONAL",
        "Credit Advice:\nEnter Custom Info -2\nNote: Header label is editable in Row 1\nOPTIONAL",
        "Credit Advice:\nEnter Custom Info -3\nNote: Header label is editable in Row 1\nOPTIONAL",
        "Credit Advice:\nEnter Custom Info -4\nNote: Header label is editable in Row 1\nOPTIONAL",
        "Credit Advice:\nEnter Custom Info -5\nNote: Header label is editable in Row 1\nOPTIONAL",
    ]

    fill_r1 = PatternFill(fill_type="solid", start_color="B4C6E7", end_color="B4C6E7")
    fill_r2 = PatternFill(fill_type="solid", start_color="D9E1F2", end_color="D9E1F2")

    font_r1 = Font(name="Calibri", size=10, bold=False, italic=False, color="000000")
    font_r2 = Font(name="Calibri", size=10, bold=False, italic=True, color="000000")
    font_data = Font(name="Calibri", size=11, bold=False, italic=False, color="000000")

    thin_side = Side(border_style="thin", color="000000")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    align_center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Render Header 1
    for c_idx, val in enumerate(headers_r1, 1):
        c = ws.cell(1, c_idx, val)
        c.fill = fill_r1
        c.font = font_r1
        c.border = thin_border
        c.alignment = align_center_wrap

    # Render Header 2
    for c_idx, val in enumerate(headers_r2, 1):
        c = ws.cell(2, c_idx, val)
        c.fill = fill_r2
        c.font = font_r2
        c.border = thin_border
        c.alignment = align_center_wrap

    # Default Transaction Date: today's date formatted as DD/MM/YYYY
    if not transaction_date_str:
        today_date = getdate(nowdate())
        transaction_date_str = today_date.strftime("%d/%m/%Y")

    curr_row = 3
    for item in payment_rows:
        beneficiary_name = str(item.get("beneficiary_name") or "").strip().upper()
        account_no = str(item.get("beneficiary_account_number") or "").strip()
        ifsc = str(item.get("ifsc") or "").strip().upper()
        tx_type = str(item.get("transaction_type") or "NEFT").strip().upper()
        debit_acc = str(item.get("debit_account_number") or debit_account_no or "10173678610").strip()
        tx_date = str(item.get("transaction_date") or transaction_date_str).strip()
        amount = flt(item.get("amount") or 0.0)
        currency = "INR"
        email = str(item.get("beneficiary_email") or "").strip() or None
        remarks = str(item.get("remarks") or "").strip() or None

        row_vals = [
            beneficiary_name,
            account_no,
            ifsc,
            tx_type,
            debit_acc,
            tx_date,
            amount,
            currency,
            email,
            remarks,
            None,  # Custom 1
            None,  # Custom 2
            None,  # Custom 3
            None,  # Custom 4
            None,  # Custom 5
        ]

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(curr_row, c_idx)
            cell.font = font_data
            if val is not None:
                cell.value = val

            # Text format (@) for bank accounts, IFSC, dates, etc.
            if c_idx in (1, 2, 3, 4, 5, 6, 8, 9, 10, 11, 12, 13, 14, 15):
                cell.number_format = "@"
            elif c_idx == 7:
                cell.number_format = "0.00"

        curr_row += 1

    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()
    filename = f"{filename_prefix}_{nowdate()}.xlsx"

    return filename, file_bytes


@frappe.whitelist()
def download_idfc_bank_batch_excel(batch_name: str):
    """
    Generates and downloads the official IDFC Bank Bulk Upload Excel (.xlsx)
    for all approved reimbursement vouchers inside a Weekly Accounts Audit Batch.
    """
    if not batch_name:
        frappe.throw(_("Batch name is required."))

    batch = frappe.get_doc("Weekly Accounts Audit Batch", batch_name)
    company = batch.company or frappe.defaults.get_user_default("Company")
    debit_account = get_company_debit_account(company)

    # Format batch transaction date (DD/MM/YYYY)
    b_date = getdate(batch.batch_date or nowdate())
    tx_date_str = b_date.strftime("%d/%m/%Y")

    items = [it for it in (batch.items or []) if it.voucher_no and it.status not in ("Disputed / Returned", "Returned to Employee", "Returned to Receptionist")]
    if not items:
        frappe.throw(_("No approved, non-disputed vouchers found in this batch."))

    # Aggregate by Employee to produce clean 1-row-per-employee bank disbursal entries
    emp_groups: Dict[str, Dict[str, Any]] = OrderedDict()

    for it in items:
        claim_name = it.voucher_no
        try:
            claim = frappe.get_doc("Employee Reimbursement Claim", claim_name)
        except Exception:
            continue

        emp_id = claim.employee or getattr(it, "employee_id", None) or getattr(it, "employee", None)
        if not emp_id:
            continue

        # Fetch latest employee bank records
        emp_info = frappe.db.get_value(
            "Employee",
            emp_id,
            ["employee_name", "bank_name", "bank_ac_no", "ifsc_code", "custom_ifsc_code", "company_email", "personal_email"],
            as_dict=True
        ) or {}

        # Prioritize claim bank account details, fallback to master
        b_name = (
            claim.get("bank_account_name")
            or claim.get("employee_name")
            or emp_info.get("employee_name")
            or emp_id
        )
        b_ac = (
            claim.get("bank_account_number")
            or emp_info.get("bank_ac_no")
            or ""
        )
        b_ifsc = (
            claim.get("bank_ifsc_code")
            or emp_info.get("custom_ifsc_code")
            or emp_info.get("ifsc_code")
            or ""
        )
        email = (
            claim.get("employee_email")
            or emp_info.get("company_email")
            or emp_info.get("personal_email")
            or ""
        )

        sanctioned_amt = flt(claim.get("sanctioned_amount") or claim.get("net_payable_amount") or claim.get("total_amount") or claim.get("total_claim_amount") or it.get("claimed_amount") or 0.0)

        if emp_id not in emp_groups:
            emp_groups[emp_id] = {
                "beneficiary_name": b_name,
                "beneficiary_account_number": str(b_ac).strip(),
                "ifsc": str(b_ifsc).strip().upper(),
                "debit_account_number": debit_account,
                "transaction_date": tx_date_str,
                "amount": sanctioned_amt,
                "currency": "INR",
                "beneficiary_email": email,
                "remarks": f"REIMBURSEMENT BATCH {batch.name}",
                "vouchers": [claim_name],
            }
        else:
            emp_groups[emp_id]["amount"] += sanctioned_amt
            emp_groups[emp_id]["vouchers"].append(claim_name)

    payment_rows = list(emp_groups.values())
    if not payment_rows:
        frappe.throw(_("Could not extract valid employee payment records from this batch."))

    filename, filebytes = generate_idfc_bank_excel_bytes(
        payment_rows=payment_rows,
        debit_account_no=debit_account,
        transaction_date_str=tx_date_str,
        filename_prefix=f"IDFC_Bank_Bulk_Upload_{batch_name}"
    )

    frappe.response["filename"] = filename
    frappe.response["filecontent"] = filebytes
    frappe.response["type"] = "binary"


@frappe.whitelist()
def download_idfc_report_excel(filters: Any = None):
    """
    Direct Excel download endpoint for the IDFC Bank Disbursal Desk Report.
    """
    if isinstance(filters, str):
        try:
            filters = json.loads(filters)
        except Exception:
            filters = {}
    elif not filters:
        filters = {}

    from ap_automation.ap_automation.report.idfc_bank_disbursal_report.idfc_bank_disbursal_report import execute

    columns, data = execute(filters)

    company = filters.get("company") or frappe.defaults.get_user_default("Company")
    debit_account = get_company_debit_account(company)
    tx_date_str = getdate(nowdate()).strftime("%d/%m/%Y")

    payment_rows = []
    for row in data:
        payment_rows.append({
            "beneficiary_name": row.get("beneficiary_name"),
            "beneficiary_account_number": row.get("beneficiary_account_number"),
            "ifsc": row.get("ifsc"),
            "debit_account_number": row.get("debit_account_number") or debit_account,
            "transaction_date": row.get("transaction_date") or tx_date_str,
            "amount": flt(row.get("amount") or 0.0),
            "currency": "INR",
            "beneficiary_email": row.get("beneficiary_email"),
            "remarks": row.get("remarks") or "IDFC REIMBURSEMENT DISBURSAL",
        })

    if not payment_rows:
        frappe.throw(_("No payment records found for the selected filters."))

    filename, filebytes = generate_idfc_bank_excel_bytes(
        payment_rows=payment_rows,
        debit_account_no=debit_account,
        transaction_date_str=tx_date_str,
        filename_prefix=f"IDFC_Bank_Report_{company or 'All'}"
    )

    frappe.response["filename"] = filename
    frappe.response["filecontent"] = filebytes
    frappe.response["type"] = "binary"


@frappe.whitelist()
def download_idfc_petty_cash_excel(voucher_name: str):
    """
    Direct endpoint: Generates and downloads the official IDFC Bank Bulk Disbursal Excel (.xlsx)
    for a Petty Cash Entry after Director Approval.
    """
    if not voucher_name:
        frappe.throw(_("Petty cash voucher name is required."))

    doc = frappe.get_doc("Petty Cash Entry", voucher_name)
    company = doc.company or frappe.defaults.get_user_default("Company")
    debit_account = get_company_debit_account(company)

    # Format transaction date (DD/MM/YYYY)
    p_date = getdate(doc.posting_date or nowdate())
    tx_date_str = p_date.strftime("%d/%m/%Y")

    # Fetch employee email and bank details if missing
    emp_info = {}
    if doc.custodian:
        emp_info = frappe.db.get_value(
            "Employee",
            doc.custodian,
            ["employee_name", "company_email", "personal_email", "bank_ac_no", "ifsc_code", "custom_ifsc_code"],
            as_dict=True
        ) or {}

    b_name = (
        doc.get("beneficiary_name")
        or emp_info.get("employee_name")
        or doc.custodian
    )
    b_ac = (
        doc.get("custodian_bank_account")
        or emp_info.get("bank_ac_no")
        or ""
    )
    b_ifsc = (
        doc.get("custodian_ifsc_code")
        or emp_info.get("custom_ifsc_code")
        or emp_info.get("ifsc_code")
        or ""
    )
    email = emp_info.get("company_email") or emp_info.get("personal_email") or ""

    payment_row = {
        "beneficiary_name": str(b_name).strip().upper(),
        "beneficiary_account_number": str(b_ac).strip(),
        "ifsc": str(b_ifsc).strip().upper(),
        "transaction_type": "NEFT",
        "debit_account_number": debit_account,
        "transaction_date": tx_date_str,
        "amount": flt(doc.total_amount or 0.0),
        "currency": "INR",
        "beneficiary_email": email,
        "remarks": f"PETTY CASH {doc.name}",
    }

    filename, filebytes = generate_idfc_bank_excel_bytes(
        payment_rows=[payment_row],
        debit_account_no=debit_account,
        transaction_date_str=tx_date_str,
        filename_prefix=f"IDFC_Bank_Petty_Cash_{doc.name}"
    )

    frappe.response["filename"] = filename
    frappe.response["filecontent"] = filebytes
    frappe.response["type"] = "binary"


@frappe.whitelist()
def download_idfc_admin_weekly_excel(voucher_name: str):
    """
    Direct endpoint: Generates and downloads the official IDFC Bank Bulk Disbursal Excel (.xlsx)
    for Admin Weekly Vendor Claim after Accounts L2 / Director Approval.
    Contains only valid vendor line items.
    """
    if not voucher_name:
        frappe.throw(_("Voucher name is required."))

    doc = frappe.get_doc("Admin Weekly Vendor Claim", voucher_name)
    company = doc.company or frappe.defaults.get_user_default("Company")
    debit_account = get_company_debit_account(company)

    # Format transaction date (DD/MM/YYYY)
    p_date = getdate(doc.posting_date or nowdate())
    tx_date_str = p_date.strftime("%d/%m/%Y")

    payment_rows = []
    for idx, row in enumerate(doc.vendor_lines or [], start=1):
        if getattr(row, "line_status", "Pending") == "Disputed / Discrepancy":
            continue

        b_name = (
            getattr(row, "bank_beneficiary_name", None)
            or getattr(row, "vendor_name", None)
            or getattr(row, "vendor", f"Vendor {idx}")
        )
        b_ac = getattr(row, "bank_account_number", "") or ""
        b_ifsc = getattr(row, "bank_ifsc_code", "") or ""
        amount = flt(getattr(row, "amount", 0.0) or 0.0)

        # Skip non-payable rows
        if amount <= 0:
            continue

        # Transaction Type resolution per IDFC Banking Rules:
        # 1. IFT: Within IDFC First Bank transfer (IFSC starts with IDFB)
        # 2. RTGS: Inter-bank transfer where amount >= 200,000 (2 Lakhs)
        # 3. NEFT: Standard Inter-bank transfer
        if b_ifsc.upper().startswith("IDFB"):
            tx_type = "IFT"
        elif amount >= 200000.0:
            tx_type = "RTGS"
        else:
            tx_type = "NEFT"

        # Fetch Supplier email if available
        supplier_email = ""
        if getattr(row, "vendor", None):
            supplier_email = frappe.db.get_value("Supplier", row.vendor, "email_id") or ""

        inv_ref = getattr(row, "invoice_number", "") or f"L{idx}"
        remarks = f"ADMIN {doc.name} {inv_ref}".strip()[:30]

        payment_rows.append({
            "beneficiary_name": str(b_name).strip().upper(),
            "beneficiary_account_number": str(b_ac).strip(),
            "ifsc": str(b_ifsc).strip().upper(),
            "transaction_type": tx_type,
            "debit_account_number": debit_account,
            "transaction_date": tx_date_str,
            "amount": amount,
            "currency": "INR",
            "beneficiary_email": supplier_email,
            "remarks": remarks,
        })

    if not payment_rows:
        frappe.throw(_("No payable vendor line items found in this weekly claim to generate IDFC bank file."))

    filename, filebytes = generate_idfc_bank_excel_bytes(
        payment_rows=payment_rows,
        debit_account_no=debit_account,
        transaction_date_str=tx_date_str,
        filename_prefix=f"IDFC_Bank_Admin_Weekly_{doc.name}"
    )

    frappe.response["filename"] = filename
    frappe.response["filecontent"] = filebytes
    frappe.response["type"] = "binary"


@frappe.whitelist()
def download_idfc_employee_reimbursement_excel(voucher_name: str):
    """
    Direct endpoint: Generates and downloads the official IDFC Bank Bulk Disbursal Excel (.xlsx)
    for Employee Reimbursement Claim after Accounts L2 / Director Approval.
    Valid for all 5 reimbursement categories.
    """
    if not voucher_name:
        frappe.throw(_("Voucher name is required."))

    doc = frappe.get_doc("Employee Reimbursement Claim", voucher_name)
    company = doc.company or frappe.defaults.get_user_default("Company")
    debit_account = get_company_debit_account(company)

    # Format transaction date (DD/MM/YYYY)
    p_date = getdate(doc.posting_date or nowdate())
    tx_date_str = p_date.strftime("%d/%m/%Y")

    # Fetch employee bank details if not present on doc
    emp_info = {}
    if doc.employee:
        emp_info = frappe.db.get_value(
            "Employee",
            doc.employee,
            ["employee_name", "company_email", "personal_email", "bank_name", "bank_ac_no", "ifsc_code", "custom_ifsc_code"],
            as_dict=True
        ) or {}

    b_name = (
        doc.get("employee_name")
        or emp_info.get("employee_name")
        or doc.employee
    )
    b_ac = (
        doc.get("bank_account_number")
        or emp_info.get("bank_ac_no")
        or ""
    )
    b_ifsc = (
        doc.get("bank_ifsc_code")
        or emp_info.get("custom_ifsc_code")
        or emp_info.get("ifsc_code")
        or ""
    )
    email = emp_info.get("company_email") or emp_info.get("personal_email") or ""

    # Net payable amount (sanctioned_amount -> net_payable_amount -> total_amount)
    amount = flt(doc.get("sanctioned_amount") or doc.get("net_payable_amount") or doc.get("total_amount") or 0.0)
    if amount <= 0:
        frappe.throw(_("No payable amount found on this reimbursement claim."))

    # Transaction Type resolution per IDFC Banking Rules:
    # 1. IFT: Within IDFC First Bank transfer (IFSC starts with IDFB)
    # 2. RTGS: Inter-bank transfer where amount >= 200,000 (2 Lakhs)
    # 3. NEFT: Standard Inter-bank transfer
    clean_ifsc = str(b_ifsc).strip().upper()
    if clean_ifsc.startswith("IDFB"):
        tx_type = "IFT"
    elif amount >= 200000.0:
        tx_type = "RTGS"
    else:
        tx_type = "NEFT"

    payment_row = {
        "beneficiary_name": str(b_name).strip().upper(),
        "beneficiary_account_number": str(b_ac).strip(),
        "ifsc": clean_ifsc,
        "transaction_type": tx_type,
        "debit_account_number": debit_account,
        "transaction_date": tx_date_str,
        "amount": amount,
        "currency": "INR",
        "beneficiary_email": email,
        "remarks": f"REIMB {doc.name}"[:30],
    }

    filename, filebytes = generate_idfc_bank_excel_bytes(
        payment_rows=[payment_row],
        debit_account_no=debit_account,
        transaction_date_str=tx_date_str,
        filename_prefix=f"IDFC_Bank_Reimbursement_{doc.name}"
    )

    frappe.response["filename"] = filename
    frappe.response["filecontent"] = filebytes
    frappe.response["type"] = "binary"

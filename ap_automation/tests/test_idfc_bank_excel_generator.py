# Copyright (c) 2026, Quanti and contributors
# For license information, please see license.txt

import io
import unittest
import openpyxl

import frappe

from ap_automation.services.idfc_bank_excel_generator import (
    generate_idfc_bank_excel_bytes,
    get_company_debit_account
)
from ap_automation.ap_automation.report.idfc_bank_disbursal_report.idfc_bank_disbursal_report import (
    execute,
    get_columns
)


class TestIDFCBankExcelGenerator(unittest.TestCase):
    def test_01_excel_structure_and_styling(self):
        sample_rows = [
            {
                "beneficiary_name": "KARTHIKEYAN M",
                "beneficiary_account_number": "046701000080723",
                "ifsc": "IOBA0000467",
                "debit_account_number": "10173678610",
                "transaction_date": "07/10/2026",
                "amount": 200.0,
                "currency": "INR",
                "beneficiary_email": "karthik@example.com",
                "remarks": "DINNER ALLOWANCE REIMBURSEMENT"
            },
            {
                "beneficiary_name": "VARAPRASADH P",
                "beneficiary_account_number": "10001979635",
                "ifsc": "IDFB0080102",
                "debit_account_number": "10173678610",
                "transaction_date": "07/10/2026",
                "amount": 800.0,
                "currency": "INR",
                "beneficiary_email": "vara@example.com",
                "remarks": "DINNER ALLOWANCE REIMBURSEMENT"
            }
        ]

        filename, file_bytes = generate_idfc_bank_excel_bytes(
            payment_rows=sample_rows,
            debit_account_no="10173678610",
            transaction_date_str="07/10/2026",
            filename_prefix="Test_IDFC_Export"
        )

        self.assertTrue(filename.endswith(".xlsx"))
        self.assertTrue(len(file_bytes) > 0)

        wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
        self.assertIn("IDFC Sheet", wb.sheetnames)
        ws = wb["IDFC Sheet"]

        # 15 Columns
        self.assertEqual(ws.max_column, 15)
        self.assertEqual(ws.max_row, 4)

        # Check Row 1 header
        self.assertEqual(ws.cell(1, 1).value, "Beneficiary Name")
        self.assertEqual(ws.cell(1, 2).value, "Beneficiary Account Number")
        self.assertEqual(ws.cell(1, 7).value, "Amount")
        self.assertEqual(ws.cell(1, 1).fill.start_color.rgb, "00B4C6E7")

        # Check Row 2 subheader
        self.assertIn("MANDATORY", ws.cell(2, 1).value)
        self.assertEqual(ws.cell(2, 1).fill.start_color.rgb, "00D9E1F2")

        # Check Row 3 data
        self.assertEqual(ws.cell(3, 1).value, "KARTHIKEYAN M")
        self.assertEqual(ws.cell(3, 2).value, "046701000080723")
        self.assertEqual(ws.cell(3, 2).number_format, "@")
        self.assertEqual(ws.cell(3, 4).value, "NEFT")
        self.assertEqual(ws.cell(3, 5).value, "10173678610")
        self.assertEqual(ws.cell(3, 6).value, "07/10/2026")
        self.assertEqual(ws.cell(3, 7).value, 200.0)
        self.assertEqual(ws.cell(3, 7).number_format, "0.00")
        self.assertEqual(ws.cell(3, 8).value, "INR")

    def test_02_debit_account_fallback(self):
        acc = get_company_debit_account(None)
        self.assertEqual(acc, "10173678610")

    def test_03_report_columns(self):
        cols = get_columns()
        self.assertEqual(len(cols), 13)
        fieldnames = [c["fieldname"] for c in cols]
        self.assertIn("beneficiary_name", fieldnames)
        self.assertIn("beneficiary_account_number", fieldnames)
        self.assertIn("ifsc", fieldnames)
        self.assertIn("amount", fieldnames)
        self.assertIn("debit_account_number", fieldnames)
